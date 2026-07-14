from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.responses import HTMLResponse, FileResponse, StreamingResponse
from pydantic import BaseModel
from typing import List
from io import BytesIO
import os
import json

import firebase_admin
from firebase_admin import credentials, firestore

app = FastAPI(title="AI Model Router API")

# Global pipeline instance to preserve learning history
from pipeline import Pipeline
from models import Request
pipeline = Pipeline()

# --- Lazy Firebase Initialization ---
# Firebase is initialized on first use to handle Vercel cold starts correctly.
_db = None
_firebase_error = None
_firebase_ready = False

def get_db():
    global _db, _firebase_error, _firebase_ready

    if _firebase_ready:
        return _db

    _firebase_ready = True

    # Try env var first (Vercel production)
    service_account_str = os.environ.get("FIREBASE_SERVICE_ACCOUNT", "").strip()
    if service_account_str:
        try:
            service_account = json.loads(service_account_str)
            if not firebase_admin._apps:
                cred = credentials.Certificate(service_account)
                firebase_admin.initialize_app(cred)
            _db = firestore.client()
            print("Firebase Admin initialized from FIREBASE_SERVICE_ACCOUNT env var!")
            return _db
        except Exception as e:
            _firebase_error = f"Env var parse/init error: {e}"
            print(_firebase_error)

    # Fallback: local JSON file (local development)
    cred_path = r"c:\Users\siddh\Downloads\ai-router-4bdff-firebase-adminsdk-fbsvc-19008248b7.json"
    if os.path.exists(cred_path):
        try:
            if not firebase_admin._apps:
                cred = credentials.Certificate(cred_path)
                firebase_admin.initialize_app(cred)
            _db = firestore.client()
            print("Firebase Admin initialized from local JSON file!")
            return _db
        except Exception as e:
            _firebase_error = f"Local file init error: {e}"
            print(_firebase_error)

    print("Firebase: no credentials found in env var or local file.")
    return None


class Message(BaseModel):
    role: str
    content: str

class ChatRequest(BaseModel):
    prompt: str
    history: List[Message] = []

@app.post("/api/chat")
async def chat(payload: ChatRequest):
    if not payload.prompt.strip():
        raise HTTPException(status_code=400, detail="Prompt cannot be empty")

    try:
        history_list = [{"role": msg.role, "content": msg.content} for msg in payload.history]
        req = Request(
            prompt=payload.prompt,
            context=json.dumps(history_list)
        )
        return StreamingResponse(
            pipeline.stream_request(req),
            media_type="text/event-stream"
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/parse-file")
async def parse_file(file: UploadFile = File(...)):
    filename = file.filename
    ext = os.path.splitext(filename)[1].lower()
    content_text = ""
    
    try:
        file_bytes = await file.read()
        file_stream = BytesIO(file_bytes)
        
        if ext == ".pptx":
            import pptx
            prs = pptx.Presentation(file_stream)
            slides_text = []
            for i, slide in enumerate(prs.slides):
                slide_parts = [f"--- Slide {i+1} ---"]
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_parts.append(shape.text.strip())
                slides_text.append("\n".join(slide_parts))
            content_text = "\n\n".join(slides_text)
            
        elif ext == ".pdf":
            import pypdf
            reader = pypdf.PdfReader(file_stream)
            pages_text = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                pages_text.append(f"--- Page {i+1} ---\n{text if text else ''}")
            content_text = "\n\n".join(pages_text)
            
        elif ext == ".docx":
            import docx
            doc = docx.Document(file_stream)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content_text = "\n".join(paragraphs)
            
        else:
            # Fallback to UTF-8 decoding for general text/code files
            try:
                content_text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                try:
                    content_text = file_bytes.decode("latin-1")
                except Exception:
                    raise HTTPException(status_code=400, detail="Unable to decode file content as text.")
                    
        return {
            "filename": filename,
            "content": content_text
        }
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Failed to parse file: {str(e)}")


# --- Firebase Auth & History Sync Endpoints ---

class AuthRequest(BaseModel):
    email: str
    password: str

@app.post("/api/auth/login")
async def auth_login(payload: AuthRequest):
    db = get_db()
    if db is None:
        raise HTTPException(
            status_code=503,
            detail=f"Firebase Admin SDK is not initialized. Error: {_firebase_error}"
        )

    email_key = payload.email.replace(".", "_")
    user_ref = db.collection("users_auth").document(email_key)
    doc = user_ref.get()

    if doc.exists:
        stored_data = doc.to_dict()
        if stored_data.get("password") != payload.password:
            raise HTTPException(status_code=400, detail="Incorrect password.")
    else:
        user_ref.set({"email": payload.email, "password": payload.password})

    return {"email": payload.email, "uid": email_key}

@app.get("/api/sessions")
async def get_sessions(email: str):
    db = get_db()
    if db is None:
        raise HTTPException(status_code=503, detail="Firebase Admin SDK is not initialized.")

    email_key = email.replace(".", "_")
    doc = db.collection("users_sessions").document(email_key).get()
    if doc.exists:
        return doc.to_dict().get("sessions", [])
    return []

class SessionsPayload(BaseModel):
    email: str
    sessions: List[dict]

@app.post("/api/sessions")
async def save_sessions(payload: SessionsPayload):
    db = get_db()
    if db is None:
        raise HTTPException(status_code=503, detail="Firebase Admin SDK is not initialized.")

    email_key = payload.email.replace(".", "_")
    db.collection("users_sessions").document(email_key).set({
        "sessions": payload.sessions
    })
    return {"status": "success"}

@app.get("/api/stats")
async def get_stats():
    stats = pipeline.learning_component.get_stats()
    from models import MODEL_CATALOG
    model_costs = {name: m.base_cost for name, m in MODEL_CATALOG.items()}

    history = pipeline.learning_component.history
    saved = 0.0
    for entry in history:
        saved += (4.40 - MODEL_CATALOG[entry.chosen_model].base_cost) * 150 / 1_000_000.0

    return {
        "total_requests": stats.get("total_requests", 0),
        "success_rate": stats.get("success_rate", 0.0),
        "total_cost": stats.get("total_estimated_cost", 0.0),
        "total_saved": max(saved, 0.0),
        "model_costs": model_costs
    }

# Debug endpoint — shows Firebase init status (safe to keep, returns no secrets)
@app.get("/api/debug")
async def debug():
    db = get_db()
    return {
        "firebase_ready": db is not None,
        "firebase_error": _firebase_error,
        "firebase_env_var_set": bool(os.environ.get("FIREBASE_SERVICE_ACCOUNT", "").strip()),
        "fireworks_key_set": bool(os.environ.get("FIREWORKS_API_KEY", "").strip())
    }

# --- Frontend Static File Serving ---
@app.get("/", response_class=HTMLResponse)
async def read_index():
    if os.path.exists("index.html"):
        return FileResponse("index.html")
    raise HTTPException(status_code=404, detail="index.html not found")

@app.get("/style.css")
async def read_css():
    if os.path.exists("style.css"):
        return FileResponse("style.css")
    raise HTTPException(status_code=404, detail="style.css not found")

@app.get("/app.js")
async def read_js():
    if os.path.exists("app.js"):
        return FileResponse("app.js")
    raise HTTPException(status_code=404, detail="app.js not found")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)
